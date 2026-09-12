from __future__ import annotations

import tempfile
from pathlib import Path


def extract_pdf(path: str | Path) -> str:
    import pymupdf4llm

    return pymupdf4llm.to_markdown(
        str(path),
        header=False,
        footer=False,
        write_images=False,
        show_progress=False,
    )


def extract_word(path: str | Path) -> str:
    from docx import Document

    document = Document(str(path))
    parts: list[str] = []
    for paragraph in document.paragraphs:
        text = (paragraph.text or "").strip()
        if not text:
            continue
        style = paragraph.style.name if paragraph.style else ""
        level = _heading_level(style)
        parts.append(f"{'#' * level} {text}" if level else text)
    for table in document.tables:
        markdown = _table_markdown(table)
        if markdown:
            parts.append(markdown)
    return "\n\n".join(parts)


def extract_text(path: str | Path) -> str:
    raw = Path(path).read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def extract_upload(uploaded) -> str:
    name = getattr(uploaded, "name", "file")
    ext = Path(name).suffix.lower()
    data = uploaded.getvalue() if hasattr(uploaded, "getvalue") else uploaded.read()
    with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as tmp:
        tmp.write(data)
        path = Path(tmp.name)
    try:
        if ext == ".pdf":
            return extract_pdf(path)
        if ext == ".docx":
            return extract_word(path)
        if ext == ".doc":
            raise ValueError(
                "Classic .doc files aren't supported. Save as .docx and upload again."
            )
        if ext == ".txt":
            return extract_text(path)
        raise ValueError("Only Word, PDF or text files are accepted.")
    finally:
        path.unlink(missing_ok=True)


def _heading_level(style_name: str) -> int | None:
    name = (style_name or "").strip()
    lowered = name.lower()
    if lowered == "title":
        return 1
    if lowered.startswith("heading"):
        rest = name[7:].strip()
        if rest.isdigit():
            return max(1, min(int(rest), 6))
        return 1
    return None


def _table_markdown(table) -> str:
    rows: list[str] = []
    for index, row in enumerate(table.rows):
        cells = [
            " ".join((cell.text or "").split())
            for cell in row.cells
        ]
        if not any(cells):
            continue
        rows.append("| " + " | ".join(cells) + " |")
        if index == 0:
            rows.append("| " + " | ".join("---" for _ in cells) + " |")
    return "\n".join(rows)
